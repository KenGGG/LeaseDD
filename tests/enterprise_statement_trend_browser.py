"""Read-only browser regression for trends backed by saved provider statement columns."""

import os
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

import openpyxl
from playwright.sync_api import sync_playwright
from leasedd.enterprise_export import export_enterprise_workbook


user = {"id": "reviewer", "username": "趋势验收", "admin": False, "csrf_token": "fixture"}
project = {
    "id": "trend-fixture", "name": "趋势测试项目", "revision": 1, "metrics": {},
    "metrics_current": False, "section_current": False, "members": [
        {"user_id": "reviewer", "username": "趋势验收", "role": "reviewer"}
    ],
}
module = {
    "module_key": "balance_sheet", "module_name": "资产负债表", "category": "statements",
    "state": "completed", "response_sha256": "a" * 64,
    "request_params": {"unit": "万元", "displayCurrency": "O", "rateType": "1"},
    "raw_payload": {},
    "parsed_payload": {"periods": ["2025年年报"] * 8, "rows": [
        {"key": "dataType", "name": "报表类型", "values": [
            "合并期末", "合并期末较年初比(%)", "合并期末同比(%)",
            "合并期末销售比(%)", "合并期末资产比(%)", "合并期末环比(%)",
            "母公司期末", "母公司期末较年初比(%)",
        ]},
        {"key": "110050", "name": "资产总计", "unit": "万元", "values": [
            "1681573.338932", "-5.576781", "-5.576781", "192.918580", "100",
            "-0.536379", "686534.494737", "13.273447",
        ]},
    ]},
}
status = {
    "source_type": "enterprise_warning", "binding": {"company_name": "趋势测试项目"},
    "import": {"id": "import-1", "state": "completed", "module_status": {}},
}


def serve(route):
    parsed = urlsplit(route.request.url)
    path = parsed.path
    assert route.request.method == "GET"
    query = parse_qs(parsed.query)
    if path.endswith("/enterprise/data") and query.get("export_format") == ["xlsx"]:
        content = export_enterprise_workbook(
            SimpleNamespace(**module), trend_key=query["trend_key"][0],
            report=query["report"][0], scopes=query["scopes"][0],
            unit=query["unit"][0], decimals=int(query["decimals"][0]),
            window_years=int(query.get("window_years", ["0"])[0]),
        )
        route.fulfill(body=content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return
    if path == "/api/me":
        value = user
    elif path == "/api/projects":
        value = [project]
    elif path == "/api/projects/trend-fixture":
        value = project
    elif path.endswith("/enterprise/data"):
        value = {"import_id": "import-1", "modules": [module]}
    elif path.endswith("/enterprise"):
        value = status
    elif path.endswith("/facts"):
        value = {"facts": [], "history": []}
    elif path.endswith("/section"):
        value = {"version": 0, "draft": None, "current": False, "input_hash": ""}
    elif path.endswith(("/documents", "/financial-statements", "/tasks", "/exports")):
        value = []
    else:
        raise AssertionError(path)
    route.fulfill(json=value)


def test_statement_trend_dialog_uses_disclosed_percentage():
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, executable_path="/usr/bin/google-chrome")
        page = browser.new_page(viewport={"width": 1600, "height": 900}, accept_downloads=True)
        page.route("**/api/**", serve)
        page.goto(os.getenv("LEASEDD_BROWSER_URL", "http://127.0.0.1:5174"))
        page.get_by_role("button", name="趋势测试项目").click()
        page.get_by_role("button", name="财务核对", exact=True).click()
        page.get_by_role("button", name="资产总计指标趋势图").click(timeout=10_000)
        dialog = page.get_by_role("dialog")
        assert "资产总计(%)" in dialog.inner_text()
        assert "-5.58" in dialog.inner_text()
        dialog.get_by_label("趋势报表类型").select_option("母公司期末")
        assert "13.27" in dialog.inner_text()
        dialog.get_by_label("趋势报表类型").select_option("合并期末")
        with page.expect_download() as event:
            dialog.get_by_role("button", name="导出Excel").click()
        workbook = openpyxl.load_workbook(BytesIO(event.value.path().read_bytes()), read_only=True, data_only=True)
        assert workbook.active.cell(1, 3).value == "资产总计(%)"
        assert workbook.active.cell(2, 3).value == -5.576781
        workbook.close()
        browser.close()


def test_finance_navigation_keeps_table_wide_without_changing_other_pages():
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, executable_path="/usr/bin/google-chrome")
        page = browser.new_page(viewport={"width": 1600, "height": 900})
        page.route("**/api/**", serve)
        page.goto(os.getenv("LEASEDD_BROWSER_URL", "http://127.0.0.1:5174"))
        page.get_by_role("button", name="趋势测试项目").click()
        page.get_by_role("button", name="财务核对", exact=True).click()
        assert page.locator(".shell").evaluate("element => element.classList.contains('finance-mode')")
        assert page.locator(".sidebar").bounding_box()["height"] < 100
        assert page.locator(".finance-content").bounding_box()["width"] > 1200
        page.get_by_role("button", name="项目概览").click()
        assert page.locator(".sidebar").bounding_box()["width"] > 200
        page.set_viewport_size({"width": 390, "height": 844})
        page.get_by_role("button", name="财务核对", exact=True).click()
        assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
        browser.close()
