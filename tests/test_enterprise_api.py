import time

import pytest
from fastapi import HTTPException
from sqlalchemy import func, select

from leasedd.app import create_app
from leasedd.contracts import EnterpriseImportRequest
from leasedd.db import Audit, Document, DocumentConversion, EnterpriseBinding, EnterpriseFinancialData, EnterpriseImport, FinancialStatement, Member, Project, Task, User
from leasedd.enterprise_warning import EnterpriseCandidate


class SearchCollector:
    def search(self, name):
        return [EnterpriseCandidate("a", "甲公司", {"symbol": "000001"}), EnterpriseCandidate("b", "甲科技", {})]


def endpoint(app, path, method):
    matches = [route.endpoint for route in app.routes if getattr(route, "path", None) == path and method in getattr(route, "methods", ())]
    assert len(matches) == 1
    return matches[0]


@pytest.fixture
def api(tmp_path):
    app = create_app(f"sqlite:///{tmp_path}/api.db", tmp_path / "files", secure_cookie=False)
    app.state.enterprise_collector = SearchCollector()
    with app.state.db.begin() as db:
        admin = User(username="admin", password_hash="x", admin=True)
        writer = User(username="writer", password_hash="x", admin=False)
        outsider = User(username="outsider", password_hash="x", admin=False)
        project = Project(name="甲公司")
        db.add_all([admin, writer, outsider, project]); db.flush()
        db.add_all([Member(project_id=project.id, user_id=admin.id, role="writer"), Member(project_id=project.id, user_id=writer.id, role="writer")])
        ids = project.id, admin.id, writer.id, outsider.id
    yield app, ids
    app.state.engine.dispose()


def test_exactly_four_routes_and_manual_search_then_binding(api):
    app, (pid, admin_id, _, _) = api
    paths = sorted(route.path for route in app.routes if "/enterprise" in getattr(route, "path", ""))
    assert paths == sorted([
        "/api/projects/{pid}/enterprise", "/api/projects/{pid}/enterprise/import",
        "/api/projects/{pid}/enterprise/retry", "/api/projects/{pid}/enterprise/data",
    ])
    import_route = endpoint(app, "/api/projects/{pid}/enterprise/import", "POST")
    with app.state.db() as db:
        admin = db.get(User, admin_id)
        result = import_route(pid, EnterpriseImportRequest(query="甲"), user=admin, db=db)
        assert [item["name"] for item in result["candidates"]] == ["甲公司", "甲科技"]
        assert db.scalar(select(func.count()).select_from(EnterpriseBinding)) == 0
        assert db.scalar(select(func.count()).select_from(Audit)) == 0
    with app.state.db() as db:
        result = import_route(pid, EnterpriseImportRequest(company_code="a", company_name="甲公司"), user=db.get(User, admin_id), db=db)
        assert result["kind"] == "enterprise_import" and result["state"] == "queued"
        assert db.scalar(select(func.count()).select_from(EnterpriseBinding)) == 1
        assert db.scalar(select(func.count()).select_from(Task)) == 1
        assert {event.action for event in db.scalars(select(Audit))} == {"bind_enterprise", "enqueue_enterprise_import"}
        duplicate = import_route(pid, EnterpriseImportRequest(company_code="a", company_name="甲公司"), user=db.get(User, admin_id), db=db)
        assert duplicate["id"] == result["id"]


def test_active_import_cannot_rebind_to_another_company(api):
    app, (pid, admin_id, _, _) = api
    route = endpoint(app, '/api/projects/{pid}/enterprise/import', 'POST')
    with app.state.db() as db:
        user = db.get(User, admin_id)
        route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=user, db=db)
        with pytest.raises(HTTPException) as error:
            route(pid, EnterpriseImportRequest(company_code='b', company_name='乙公司'), user=user, db=db)
        assert error.value.status_code == 409
        assert db.scalar(select(EnterpriseBinding)).company_code == 'a'


def test_new_queued_import_is_visible_in_status(api):
    app, (pid, admin_id, _, _) = api
    route = endpoint(app, '/api/projects/{pid}/enterprise/import', 'POST')
    status = endpoint(app, '/api/projects/{pid}/enterprise', 'GET')
    with app.state.db() as db:
        user = db.get(User, admin_id)
        route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=user, db=db)
        task = db.scalar(select(Task));task.state = 'completed'
        record = db.scalar(select(EnterpriseImport));record.state = 'completed';record.completed_at = time.time()
        db.commit()
        newest = route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=user, db=db)
        assert newest['state'] == 'queued'
        assert status(pid, user=user, db=db)['import']['state'] == 'queued'


def test_project_reviewer_can_refresh_bound_company_without_rebinding(api):
    app, (pid, admin_id, member_id, _) = api
    route = endpoint(app, '/api/projects/{pid}/enterprise/import', 'POST')
    with app.state.db.begin() as db:
        member = db.scalar(select(Member).where(Member.project_id == pid, Member.user_id == member_id))
        member.role = 'reviewer'
    with app.state.db() as db:
        route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=db.get(User, admin_id), db=db)
    with app.state.db.begin() as db:
        task = db.scalar(select(Task).where(Task.project_id == pid))
        task.state = 'completed'
    with app.state.db() as db:
        result = route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=db.get(User, member_id), db=db)
        assert result['state'] == 'queued'
        assert db.scalar(select(func.count()).select_from(Task)) == 2
        assert db.scalar(select(EnterpriseBinding)).company_code == 'a'


def test_project_member_refresh_cannot_search_or_change_binding(api):
    app, (pid, admin_id, member_id, outsider_id) = api
    route = endpoint(app, '/api/projects/{pid}/enterprise/import', 'POST')
    with app.state.db() as db:
        route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=db.get(User, admin_id), db=db)
    with app.state.db.begin() as db:
        db.scalar(select(Task).where(Task.project_id == pid)).state = 'completed'
    with app.state.db() as db:
        member = db.get(User, member_id)
        for request in (EnterpriseImportRequest(query='乙'),
                        EnterpriseImportRequest(company_code='b', company_name='乙公司'),
                        EnterpriseImportRequest(company_code='a', company_name='冒名公司')):
            with pytest.raises(HTTPException) as denied:
                route(pid, request, user=member, db=db)
            assert denied.value.status_code == 403
        with pytest.raises(HTTPException) as denied:
            route(pid, EnterpriseImportRequest(company_code='a', company_name='甲公司'), user=db.get(User, outsider_id), db=db)
        assert denied.value.status_code == 403
        assert db.scalar(select(EnterpriseBinding)).company_name == '甲公司'


def test_rebinding_does_not_relabel_previous_company_financial_data(api):
    app,(pid,admin_id,_,_)=api
    route=endpoint(app,'/api/projects/{pid}/enterprise/import','POST')
    data=endpoint(app,'/api/projects/{pid}/enterprise/data','GET')
    with app.state.db() as db:
        user=db.get(User,admin_id)
        route(pid,EnterpriseImportRequest(company_code='a',company_name='甲公司'),user=user,db=db)
        task=db.scalar(select(Task));task.state='completed'
        record=db.scalar(select(EnterpriseImport));record.state='completed';record.completed_at=time.time()
        db.commit()
        route(pid,EnterpriseImportRequest(company_code='b',company_name='乙公司'),user=user,db=db)
        assert data(pid,user=user,db=db)['import_id'] is None


def test_permissions_status_data_and_retry_audit(api):
    app, (pid, admin_id, writer_id, outsider_id) = api
    import_route = endpoint(app, "/api/projects/{pid}/enterprise/import", "POST")
    status_route = endpoint(app, "/api/projects/{pid}/enterprise", "GET")
    data_route = endpoint(app, "/api/projects/{pid}/enterprise/data", "GET")
    retry_route = endpoint(app, "/api/projects/{pid}/enterprise/retry", "POST")
    with app.state.db() as db:
        with pytest.raises(HTTPException) as denied:
            import_route(pid, EnterpriseImportRequest(query="甲"), user=db.get(User, writer_id), db=db)
        assert denied.value.status_code == 403
        with pytest.raises(HTTPException) as denied:
            status_route(pid, user=db.get(User, outsider_id), db=db)
        assert denied.value.status_code == 403
        assert status_route(pid, user=db.get(User, writer_id), db=db)["binding"] is None
        assert data_route(pid, category=None, module_key=None, user=db.get(User, writer_id), db=db)["modules"] == []
        import_route(pid, EnterpriseImportRequest(company_code="a", company_name="甲公司"), user=db.get(User, admin_id), db=db)
    with app.state.db.begin() as db:
        task = db.scalar(select(Task).where(Task.project_id == pid)); task.state = "completed"
        task.result = {**task.result, "failed_modules": ["module-5"]}
    with app.state.db() as db:
        retry = retry_route(pid, user=db.get(User, admin_id), db=db)
        assert retry["kind"] == "enterprise_import"
        assert "retry_enterprise_import" in {event.action for event in db.scalars(select(Audit))}


def test_readable_enterprise_import_replaces_pdf_statement_view(api):
    app, (pid, admin_id, _, _) = api
    import_route = endpoint(app, "/api/projects/{pid}/enterprise/import", "POST")
    statement_route = endpoint(app, "/api/projects/{pid}/financial-statements", "GET")
    with app.state.db() as db:
        import_route(pid, EnterpriseImportRequest(company_code="a", company_name="甲公司"), user=db.get(User, admin_id), db=db)
    with app.state.db.begin() as db:
        task = db.scalar(select(Task).where(Task.project_id == pid))
        record = db.scalar(select(EnterpriseImport).where(EnterpriseImport.project_id == pid))
        record.state = "completed"; record.completed_at = time.time()
        db.add(EnterpriseFinancialData(import_id=record.id, category="statements", module_key="balance", module_name="资产负债表", module_order=0,
            endpoint_path="/getThreeReports", request_params={"statement_type":"balance_sheet","unit":"元"}, raw_payload={"data":{}},
            parsed_payload={"periods":["2025-12-31"],"rows":[{"name":"资产总计","key":"assets","values":["100"]}]},
            response_sha256="a"*64, state="completed", error=None, collected_at=time.time()))
        document = Document(project_id=pid,name="old.pdf",sha256="b"*64,path="old.pdf",model_allowed=False,parse_state="text_available",created_by=admin_id)
        db.add(document); db.flush()
        conversion = DocumentConversion(project_id=pid,document_id=document.id,original_sha256=document.sha256,tool="mineru",tool_version="1",state="completed",markdown_path="old.md",markdown_sha256="c"*64,error=None,created_at=time.time(),completed_at=time.time())
        db.add(conversion); db.flush()
        db.add(FinancialStatement(project_id=pid,document_id=document.id,conversion_id=conversion.id,run_id=None,statement_type="income_statement",entity="旧PDF",scope="consolidated",period="2024",period_normalized="2024",period_kind="year",currency="CNY",raw_unit="元",unit_scale="1",source_start_line=1,source_end_line=2,state="extracted",issues=[],created_at=time.time()))
    with app.state.db() as db:
        views = statement_route(pid, run_id="", user=db.get(User, admin_id), db=db)
        assert len(views) == 1
        assert views[0]["source_type"] == "enterprise_warning"
        assert views[0]["statement_type"] == "balance_sheet"


def test_new_partial_import_does_not_replace_previous_complete_data(api):
    app, (pid, admin_id, _, _) = api
    data_route = endpoint(app, "/api/projects/{pid}/enterprise/data", "GET")
    status_route = endpoint(app, "/api/projects/{pid}/enterprise", "GET")
    with app.state.db.begin() as db:
        tasks = [Task(project_id=pid,kind="enterprise_import",mode="qyyjt",input_revision=0,input_hash="a"*64,created_by=admin_id,created_at=n) for n in (1,3)]
        db.add_all(tasks); db.flush()
        complete = EnterpriseImport(project_id=pid, task_id=tasks[0].id, state="completed", quality_state="passed", module_status={}, started_at=1, completed_at=2)
        partial = EnterpriseImport(project_id=pid, task_id=tasks[1].id, state="partial", quality_state="passed_with_gaps", module_status={}, started_at=3, completed_at=4)
        db.add_all([complete, partial]); db.flush()
        complete_id, partial_id = complete.id, partial.id
    with app.state.db() as db:
        assert data_route(pid, category=None, module_key=None, user=db.get(User, admin_id), db=db)["import_id"] == complete_id
        assert status_route(pid, user=db.get(User, admin_id), db=db)["import"]["id"] == partial_id


def test_excel_export_reuses_data_route_and_project_authorization(api):
    from io import BytesIO
    import openpyxl
    app, (pid, admin_id, _, outsider_id) = api
    import_route=endpoint(app,'/api/projects/{pid}/enterprise/import','POST')
    data_route=endpoint(app,'/api/projects/{pid}/enterprise/data','GET')
    with app.state.db() as db:
        import_route(pid,EnterpriseImportRequest(company_code='a',company_name='甲公司'),user=db.get(User,admin_id),db=db)
    with app.state.db.begin() as db:
        record=db.scalar(select(EnterpriseImport).where(EnterpriseImport.project_id==pid))
        record.state='completed';record.completed_at=time.time()
        db.add(EnterpriseFinancialData(import_id=record.id,category='indicators',module_key='main_indicators',module_name='主要财务指标',module_order=0,
            endpoint_path='/getMainIndicators',request_params={},raw_payload={'original':'unchanged'},
            parsed_payload={'periods':['2025年年报'],'rows':[{'name':'营业收入','key':'220006','unit':'万元','values':['100.25']},{'name':'报表类型','key':'dataType','values':['合并期末']}]},
            response_sha256='a'*64,state='completed',collected_at=time.time()))
    with app.state.db() as db:
        response=data_route(pid,module_key='main_indicators',export_format='xlsx',user=db.get(User,admin_id),db=db)
        assert response.media_type=='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        book=openpyxl.load_workbook(BytesIO(response.body));assert book.active['C3'].value==100.25;book.close()
        assert '.xlsx' in response.headers['content-disposition']
        assert db.scalar(select(EnterpriseFinancialData)).raw_payload=={'original':'unchanged'}
        trend=data_route(pid,module_key='main_indicators',export_format='xlsx',trend_key='220006',report='annual',scopes='合并期末',user=db.get(User,admin_id),db=db)
        book=openpyxl.load_workbook(BytesIO(trend.body))
        assert list(book.active.values)==[('序号','报告期','营业收入（万元）'),('1','2025年年报',100.25)]
        book.close()
        with pytest.raises(HTTPException) as denied:
            data_route(pid,module_key='main_indicators',export_format='xlsx',trend_key='220006',report='annual',scopes='合并期末',user=db.get(User,outsider_id),db=db)
        assert denied.value.status_code==403
        with pytest.raises(HTTPException) as stale:
            data_route(pid,module_key='main_indicators',export_format='xlsx',expected_hash='b'*64,user=db.get(User,admin_id),db=db)
        assert stale.value.status_code==409


def test_excel_request_without_readable_import_is_not_a_json_file_disguised_as_excel(api):
    app,(pid,admin_id,_,_)=api
    data_route=endpoint(app,'/api/projects/{pid}/enterprise/data','GET')
    with app.state.db() as db:
        with pytest.raises(HTTPException) as missing:
            data_route(pid,module_key='main_indicators',export_format='xlsx',user=db.get(User,admin_id),db=db)
        assert missing.value.status_code==404


def test_module_index_skips_large_json_and_detail_is_pinned_to_readable_import(api):
    from sqlalchemy import event
    app,(pid,admin_id,_,outsider_id)=api
    create=endpoint(app,'/api/projects/{pid}/enterprise/import','POST')
    data=endpoint(app,'/api/projects/{pid}/enterprise/data','GET')
    with app.state.db() as db:
        create(pid,EnterpriseImportRequest(company_code='a',company_name='甲公司'),user=db.get(User,admin_id),db=db)
    with app.state.db.begin() as db:
        record=db.scalar(select(EnterpriseImport));record.state='completed';record.completed_at=time.time();record_id=record.id
        for index,key in enumerate(['balance_sheet','income_statement']):
            db.add(EnterpriseFinancialData(import_id=record.id,category='statements',module_key=key,module_name=key,module_order=index,
                endpoint_path='/reports',request_params={'code':'a'},raw_payload={'large':'x'*100000},parsed_payload={'rows':[]},
                response_sha256='a'*64,state='completed',collected_at=time.time()))
        # Exercise legacy company identity lookup too: it must not fetch payloads.
        task=db.get(Task,record.task_id);task.result={}
    statements=[]
    def capture(connection,cursor,statement,parameters,context,many):statements.append(statement)
    with app.state.db() as db:
        user=db.get(User,admin_id)
        event.listen(app.state.engine,'before_cursor_execute',capture)
        try:summary=data(pid,summary_only=True,user=user,db=db)
        finally:event.remove(app.state.engine,'before_cursor_execute',capture)
        assert summary['import_id']==record_id and len(summary['modules'])==2
        assert all('raw_payload' not in m and 'parsed_payload' not in m for m in summary['modules'])
        assert not any('raw_payload' in sql or 'parsed_payload' in sql for sql in statements)
        detail=data(pid,module_key='balance_sheet',import_id=record_id,user=user,db=db)
        assert [m['module_key'] for m in detail['modules']]==['balance_sheet']
        assert detail['modules'][0]['raw_payload']=={'large':'x'*100000}
        with pytest.raises(HTTPException) as changed:
            data(pid,module_key='balance_sheet',import_id=record_id,expected_hash='b'*64,user=user,db=db)
        assert changed.value.status_code==409
        with pytest.raises(HTTPException) as stale:data(pid,module_key='balance_sheet',import_id='old',user=user,db=db)
        assert stale.value.status_code==409
        with pytest.raises(HTTPException) as denied:data(pid,summary_only=True,user=db.get(User,outsider_id),db=db)
        assert denied.value.status_code==403
