"""Build a reviewable workbook and standalone filterable report from captured cells."""
import json
from pathlib import Path
from collections import Counter
from html import escape
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

out=Path('runtime/acceptance/full-reconciliation')
rows=json.loads((out/'all-cells.json').read_text())
coverage=json.loads((out/'coverage.json').read_text())
live=json.loads((out/'live-before.json').read_text())
counts=Counter(r['结论'] for r in rows)
intro=[
 ['项目','铭普光磁 · 本地 PDF 与企业预警通逐项核对'],
 ['结论','本轮修复已核验；完整一致性验收尚未通过。数值一致不等于所有样式、行序和计算口径一致。'],
 ['本地资料','2026年半年报、2025年报、2024年报、2023年报，共4份已上传PDF；原件摘要和Markdown摘要见资料清单。'],
 ['对账范围','合并口径；参考界面已取得有效表格的三张报表、主要指标、7类分析、17类附注；逐格比对2022年以来捕获的列。'],
 ['范围限制','参考页面更早历史列未纳入。未能打开独立数据页的14个菜单另列为未核验，不将此前残留页面当作目标栏目。'],
 ['数据来源','本地数值仅来自已上传PDF的转换表格、已有来源核验值及确定性计算。参考数据只用于验收，没有回填项目或调用Agnes补数。'],
 ['年初与年末','资产负债表2023-01-01归入2022-12-31比较列；原始日期、候选、来源仍保留。有重述或不同金额时显示冲突，不能假定必与2022原年报一致。'],
 ['一致含义','同期间、同口径下按参考显示单位及小数精度相同，不代表原始数值每一位相同或已独立人工审核。'],
 ['本地缺失含义','目标单元格未匹配到值：包括未上传对应报告、字段尚未实现、分类或行名未完全映射。原因列分开列明，不能据此认定PDF没有数值。'],
 ['本地多出含义','参考显示为空，本地原报告有数值；保留原披露，不为了页面相同删数。'],
 ['来源定位','PDF文件名 + 经摘要校验的Markdown行号；不是PDF物理页码。原始参考截图和表格存于reference目录。'],
 ['统计限制','包含未上传季度报告的列、空值和未映射字段，不能将总体一致格数换算成提取准确率。'],
 ['当前展示限制','三张报表仍保留原报告科目顺序；部分附注行列、毛利率/成本构成、历史名单和计算口径尚未完全复刻。'],
 ['参考页面','https://www.qyyjt.cn/detail/enterprise/financialStatements?code=5C36774FC05682B8AD4996EA6B0325CB&type=company'],
]
issues=[
 {'事项':'年初余额重复列','处理':'已修复','证据与结论':'2023-01-01映射2022-12-31；剔除只有补充指标的空报表列；原表日期保留于来源窗口。','后续':'不同金额不静默覆盖，仍按来源冲突处理。'},
 {'事项':'附注中途失去合并口径','处理':'已修复','证据与结论':'“七、合并财务报表项目注释”内的“二、联营企业”不再误判为顶层章节。','后续':'已有材料刷新生效，无需再次上传。'},
 {'事项':'PDF断行数字截断','处理':'已修复并回归','证据与结论':'509,352,8 + 17.46 → 509,352,817.46；30,000,514.4 + 6 → 30,000,514.46。仅在相邻行、数字语法能证明续接时拼接。','后续':'其他不完整数值留空，不自动猜补。'},
 {'事项':'附注新增明细','处理':'已补充','证据与结论':'客户、供应商、账龄、前五名、重大坏账、受限资产、财务费用、非经常性损益、境内审计费用等已接入原文证据。','后续':'部分历史明细没有对应报告，仍为空。'},
 {'事项':'2025应收账龄合计','处理':'口径不同，未强行改值','证据与结论':'参考2025合计采用净值，早期年度采用余额；本地按原报告账龄余额显示。','后续':'需要统一明确余额/坏账准备/净值三行映射，避免净值与余额混用。'},
 {'事项':'前五名应收款比例','处理':'保留原披露','证据与结论':'部分参考期间比例分母为纯应收账款，原报告披露比例包含合同资产。','后续':'需明确参考平台分母规则；不能把重算结果冒充原报告比例。'},
 {'事项':'2025非经常性损益分类','处理':'保留原披露','证据与结论':'PDF单列资金占用费73,540.74元，其他为-7,115,688.61元；参考资金占用费为空、其他为-7,042,147.87元。','后续':'如果合并展示，应同时保留原分类和变换规则。'},
 {'事项':'2023每股指标','处理':'需人工核实分母','证据与结论':'2023股份变动表期末212,271,108股，股本附注211,520,000股；参考每股指标采用后一分母。','后续':'不能只为匹配参考值替换来源；需核实股份变动表与股本口径。'},
 {'事项':'2022受限资产','处理':'需补原报告核实','证据与结论':'2023年报上年比较列合计493,159,435.31元，与参考2022年末不一致。','后续':'需2022原年报，区分重述、项目范围和数据源差异。'},
 {'事项':'主营构成','处理':'未完全完成','证据与结论':'已展示收入和已披露成本；细分总计、部分产品成本及毛利/毛利率尚未完全映射。','后续':'按收入、成本、毛利、毛利率与产品/行业/地区四类结构逐项完善。'},
 {'事项':'三表与分析','处理':'仍有展示与公式缺口','证据与结论':'部分参考科目是附注补充项或合并项；本地目前主要按原表科目展示。部分现金/债务/每股公式与参考口径不同。','后续':'按差异表补字段和可追溯公式，不能猜测、补零或回填外部结果。'},
 {'事项':'参考菜单未打开','处理':'未完成核验','证据与结论':'部分重要逾期款项、投资性房地产、长期应收款等菜单未显示独立目标数据页，见栏目覆盖。','后续':'不能宣称已完全复制，也不能把上一页内容当作目标数据。'},
]
style=[
 {'元素':'左侧财务栏目','当前':'三表、主要指标、7类分析、财务附注和分组子菜单','状态':'已实现主要层级','剩余':'部分参考菜单未取得有效页面，不能认定逐项一致'},
 {'元素':'年初/年末列','当前':'2023年初归入2022年末；不再产生重复比较列','状态':'已修复并浏览器验证','剩余':'会计重述差异必须保留'},
 {'元素':'附注横向年度比较','当前':'项目名固定，最新期间及历史年报横向排列','状态':'已实现','剩余':'个别子行数量、单位设置与参考仍有差异'},
 {'元素':'客户/供应商/前五名名单','当前':'按年度分段，名称、金额、比例分列','状态':'已实现','剩余':'未上传报告的历史名单、匿名客户跨年增长不推断'},
 {'元素':'账龄/坏账/净值','当前':'按账龄列余额，未知坏账明细留空','状态':'部分完成','剩余':'参考不同年份总计口径和坏账子行还需逐期适配'},
 {'元素':'主营构成','当前':'收入和成本分产品、行业、地区','状态':'部分完成','剩余':'毛利率、维度总计、部分成本未完成'},
 {'元素':'三张报表科目','当前':'以原表科目顺序、单位和来源展示','状态':'尚非逐行复刻','剩余':'参考的组合行、附注补充行和公告日期等尚未齐全'},
 {'元素':'数字清晰度','当前':'统一深色数字、等宽数字间距、右对齐；可横向滚动','状态':'浏览器验收通过','剩余':'不是像素级一致承诺'},
 {'元素':'来源查看','当前':'点击数字查看PDF转换原表、原文行号与原件下载','状态':'已实现','剩余':'新增本地可追溯能力，默认不铺满页面'},
 {'元素':'CSV','当前':'金额与百分比单独格式化，名单和横向表均可导出','状态':'浏览器验收通过','剩余':'空白/来源冲突不能作为零'},
]
summary=[{'栏目':cat,**{s:sum(r['栏目']==cat and r['结论']==s for r in rows) for s in ['一致','双方空值','本地缺失','本地多出','数值差异']}} for cat in dict.fromkeys(r['栏目'] for r in rows)]
docs=[{'原件':d['name'],'原件SHA256':d['original_sha256'],'转换MarkdownSHA256':d['markdown_sha256'],'识别状态':d['parse_state']} for d in live['documents']]
wb=Workbook();wb.remove(wb.active)
def sheet(title,entries):
 ws=wb.create_sheet(title)
 if isinstance(entries[0],dict):ws.append(list(entries[0]));entries=[list(r.values()) for r in entries]
 else:ws.append(['说明项','内容'])
 for row in entries:ws.append(row)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.font=Font(color='FFFFFF',bold=True,size=11);c.fill=PatternFill('solid',fgColor='173E66')
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True);cell.font=Font(name='Microsoft YaHei',size=11,color='18283B')
   if cell.value is not None:cell.data_type='s' if isinstance(cell.value,str) else cell.data_type
  if row[0].row%2==0:
   for cell in row:cell.fill=PatternFill('solid',fgColor='F3F6FA')
 for i in range(1,ws.max_column+1):
  heading=str(ws.cell(1,i).value);ws.column_dimensions[get_column_letter(i)].width=60 if heading in ('内容','PDF来源','原因','证据与结论','剩余','后续') else 36 if heading in ('项目','原件','原件SHA256','转换MarkdownSHA256','当前') else 24
 for row in ws.iter_rows(min_row=2):ws.row_dimensions[row[0].row].height=50 if title in ('阅读说明','关键问题','样式比对') else 32
sheet('阅读说明',intro);sheet('关键问题',issues);sheet('栏目覆盖',coverage);sheet('样式比对',style);sheet('汇总',summary)
sheet('全部目标单元格',rows);sheet('待处理差异',[r for r in rows if r['结论'] not in ('一致','双方空值')]);sheet('资料清单',docs)
wb.save(out/'铭普光磁_财务逐项比对.xlsx')
check=load_workbook(out/'铭普光磁_财务逐项比对.xlsx',read_only=True)
assert check['全部目标单元格'].max_row==len(rows)+1
check.close()
payload=json.dumps(rows,ensure_ascii=False).replace('<','\\u003c')
options=''.join('<option>'+escape(cat)+'</option>' for cat in dict.fromkeys(r['栏目'] for r in rows))
headers=''.join('<th>'+escape(k)+'</th>' for k in rows[0])
introhtml=''.join('<p><b>'+escape(k)+'：</b>'+escape(v)+'</p>' for k,v in intro[1:])
html='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>铭普光磁财务逐项比对</title><style>body{margin:28px;background:#f3f6fa;color:#172b41;font:15px/1.65 system-ui,sans-serif}h1{font-size:25px}section,.filters{background:white;padding:18px;margin:16px 0;border:1px solid #c6d3e0;border-radius:6px}select,input{padding:8px;font:inherit;border:1px solid #8195ab;margin:4px}a{color:#1655a3}table{border-collapse:collapse;background:white;width:100%;font-variant-numeric:tabular-nums}th,td{border:1px solid #ccd6e0;padding:9px;vertical-align:top;min-width:90px}th{background:#173e66;color:white;position:sticky;top:0}td:nth-child(4),td:nth-child(5){text-align:right;white-space:nowrap}tr:nth-child(even){background:#f6f8fb}td:nth-child(7),td:nth-child(8){min-width:260px;max-width:420px}#wrap{overflow:auto;max-height:75vh}.diff{color:#9a3412}small{color:#4b6074}</style><h1>铭普光磁 · 财务逐项比对</h1><p><b>完整一致性验收尚未通过。</b>本页逐格展示已捕获参考值、本地值、原因与来源。</p><a href="铭普光磁_财务逐项比对.xlsx">下载完整 Excel（含样式比对和栏目覆盖）</a><section><details><summary>核对范围与使用说明</summary>INTRO</details></section><div class="filters"><label>栏目 <select id="category"><option value="">全部</option>OPTIONS</select></label><label>结果 <select id="status"><option value="">全部</option><option>一致</option><option>双方空值</option><option>本地缺失</option><option>本地多出</option><option>数值差异</option></select></label><label>搜索 <input id="query" placeholder="科目、期间、原因"></label><p id="count"></p></div><div id="wrap"><table><thead><tr>HEADERS</tr></thead><tbody id="body"></tbody></table></div><script>const rows=PAYLOAD;const category=document.querySelector('#category'),status=document.querySelector('#status'),query=document.querySelector('#query'),body=document.querySelector('#body');function render(){const chosen=rows.filter(r=>(!category.value||r['栏目']===category.value)&&(!status.value||r['结论']===status.value)&&(!query.value||Object.values(r).join(' ').includes(query.value)));body.replaceChildren();for(const r of chosen){const tr=document.createElement('tr');for(const [k,v]of Object.entries(r)){const td=document.createElement('td');td.textContent=v;if(k==='结论'&&!['一致','双方空值'].includes(v))td.className='diff';tr.append(td)}body.append(tr)}document.querySelector('#count').textContent=`显示 ${chosen.length} / ${rows.length} 格。统计包含缺报告与空值，不能作为提取准确率。`}for(const e of [category,status,query])e.addEventListener('input',render);render();</script></html>'''.replace('INTRO',introhtml).replace('OPTIONS',options).replace('HEADERS',headers).replace('PAYLOAD',payload)
(out/'铭普光磁_财务逐项比对.html').write_text(html)
print(json.dumps({'cells':len(rows),'counts':counts,'verified_categories':len(summary),'coverage_entries':len(coverage),'files':['铭普光磁_财务逐项比对.xlsx','铭普光磁_财务逐项比对.html']},ensure_ascii=False))
