"""Excel presentation of a saved enterprise module; no collection or mutation."""
from __future__ import annotations

from decimal import Decimal, localcontext
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .finance_extract import UNIT_SCALES


def _blank(value):
    return value is None or str(value).strip() in {'', '-', '--', '—'}


def _period(value):
    value = str(value)
    match = re.fullmatch(r'(\d{4})(0331|0630|0930|1231)', value)
    return match[1]+'年'+{'0331':'一季报','0630':'中报','0930':'三季报','1231':'年报'}[match[2]] if match else value


def _kind(period):
    for key, pattern in [('annual',r'年报|1231|12-31'),('half',r'中报|半年|0630|06-30'),('q3',r'三季报|0930|09-30'),('q1',r'一季报|0331|03-31')]:
        if re.search(pattern,period):return key
    return ''


def _order(period):
    year=re.match(r'\d{4}',period)
    return int(year[0])*10+{'annual':4,'q3':3,'half':2,'q1':1}.get(_kind(period),0) if year else 0


def _amount(value, source_unit, target_unit, decimals):
    if value is None:return ''
    text=str(value)
    if source_unit not in UNIT_SCALES and source_unit not in {'%','倍','天','元/股'}:return text
    if not re.fullmatch(r'-?\d+(?:\.\d+)?(?:[eE][+-]?\d{1,3})?',text):return text
    with localcontext() as context:
        context.prec=max(50,len(text)+20)
        number=Decimal(text)
        if source_unit in UNIT_SCALES:number=number*UNIT_SCALES[source_unit]/UNIT_SCALES[target_unit]
        if number==0:number=abs(number)
        return number


def _flatten(rows, source_rows):
    for row in rows:
        key=row.get('key',row.get('value'))
        values=row.get('values')
        if values is None:values=[source.get(key) for source in source_rows]
        yield {**row,'values':values}
        yield from _flatten(row.get('children') or [],source_rows)


def select_currency_variant(module,currency='',rate=''):
    from types import SimpleNamespace
    if not currency and not rate:return module
    currency=currency or 'O';rate=rate or '1'
    variants=(module.parsed_payload or {}).get('variants') or []
    if not variants:
        if currency=='O' and rate=='1':return module
        raise ValueError('currency_variant_unavailable')
    for index,variant in enumerate(variants):
        params=variant['request_params']
        if params.get('displayCurrency')==currency and str(params.get('rateType'))==rate:
            responses=(module.raw_payload or {}).get('responses') or []
            if index>=len(responses):raise ValueError('currency_variant_unavailable')
            return SimpleNamespace(module_name=module.module_name,module_key=module.module_key,category=module.category,
                request_params={**module.request_params,**params},raw_payload=responses[index]['payload'],
                parsed_payload=variant['parsed'],response_sha256=module.response_sha256)
    raise ValueError('currency_variant_unavailable')


def _column_selected(value,scopes,data_kinds):
    text=str(value or '');match=re.match(r'^(合并|母公司)期[初末]',text)
    scope=match[0] if match else text;suffix=text[len(scope):]
    kind={'较年初比(%)':'较年初增长率','同比(%)':'同比增长率','销售比(%)':'销售百分比','资产比(%)':'资产百分比','环比(%)':'环比增长率'}.get(suffix,suffix or '原始报表')
    if text in {'原始报表','同比','占收入比'}:scope,kind='',text
    return (not scopes or scopes=='all' or scope in scopes.split(',')) and (not data_kinds or data_kinds=='all' or kind in data_kinds.split(','))


def _statement_trend_output(module, periods, rows, key, report, start, end, window_years, scope, unit, decimals):
    amount=next(row for row in rows if row.get('key')==key)
    metadata=(module.parsed_payload or {}).get('metadata') or {}
    export_headers=metadata.get('headExport') or []
    row_index=rows.index(amount)
    source_unit=amount.get('unit') or module.request_params.get('unit') or '万元'
    if row_index+1<len(export_headers):
        match=re.search(r'[（(]([^（）()]+)[）)]$',str(export_headers[row_index+1]))
        if match:source_unit=match[1]
    data_types=next((row.get('values') or [] for row in rows if row.get('key')=='dataType'),[])
    currency_values=next((row.get('values') or [] for row in rows if row.get('key')=='displayCurrency'),[])
    chosen=[period for period in dict.fromkeys(periods) if (not start or period[:4]>=start)
            and (not end or period[:4]<=end) and _kind(period)==report]
    if window_years:
        quarter=lambda period:(_order(period)//10)*4+_order(period)%10
        cutoff=max(map(quarter,periods),default=0)-window_years*4
        chosen=[period for period in chosen if quarter(period)>cutoff]
    chosen.sort(key=_order,reverse=True)
    suffixes={'balance_sheet':[('较年初比(%)','%',""),('同比(%)','%','同比增长率'),('销售比(%)','%','销售百分比'),('资产比(%)','%','资产百分比'),('环比(%)','%','环比增长率')],
              'income_statement':[('',source_unit,""),('同比(%)','%','同比增长率'),('销售比(%)','%','销售百分比')],
              'cash_flow_statement':[('',source_unit,""),('同比(%)','%','同比增长率')]}[module.module_key]
    positions=[[next((index for index,period in enumerate(periods)
                      if period==selected and index<len(data_types) and data_types[index]==scope+suffix),None)
                for suffix,_,_ in suffixes] for selected in chosen]
    currencies=[str(currency_values[indices[0]]) if indices[0] is not None and indices[0]<len(currency_values)
                and not _blank(currency_values[indices[0]]) else '' for indices in positions]
    unique=set(currencies)
    currency=currencies[0] if len(unique)==1 and currencies else ''
    name=str(amount.get('name') or amount.get('label') or key)
    first_label=f'{name}（{unit}{currency}）' if suffixes[0][1] in UNIT_SCALES else f'{name}(%)'
    header=['序号','报告期',first_label,*[label+'(%)' for _,_,label in suffixes[1:]]]
    if len(unique)>1:header.append('币种')
    output=[header]
    values=amount.get('values') or []
    for number,(period,indices) in enumerate(zip(chosen,positions),1):
        row=[number,period]
        for index,(_,source,_) in zip(indices,suffixes):
            raw=values[index] if index is not None and index<len(values) else None
            row.append(_amount(raw,source,unit,decimals))
        if len(unique)>1:row.append(currencies[number-1])
        output.append(row)
    return output


def _workbook_bytes(module, output, source_style, decimals, *, source_label='指标名称', row_numbers=None, cell_formats=None):
    if source_style:
        numbers=row_numbers if row_numbers is not None else range(1,len(output))
        output=[['数据来源：企业预警通'],['序号',source_label,*output[0][1:]],*[[number,*row] for number,row in zip(numbers,output[1:])]]
    book=Workbook();sheet=book.active
    sheet.title=re.sub(r'[\\/*?:\[\]]','',module.module_name or '财务数据')[:31] or '财务数据'
    sheet.freeze_panes='C3' if source_style else 'B2'
    for row_index,row in enumerate(output,1):
        for column,value in enumerate(row,1):
            cell=sheet.cell(row_index,column)
            if isinstance(value,Decimal) and len(''.join(map(str,value.as_tuple().digits)).rstrip('0'))<=15:
                cell.value=value;cell.number_format=(cell_formats or {}).get((row_index,column),'#,##0'+('.'+'0'*decimals if decimals else ''))
            elif source_style and column==1 and row_index>2:
                cell.value=value;cell.number_format='General' if isinstance(value,str) else '0'
            else:
                cell.value=format(value,f',.{max(decimals,-value.as_tuple().exponent)}f') if isinstance(value,Decimal) else '' if value is None else str(value)
                cell.data_type='s'
            cell.font=Font(name='Microsoft YaHei',size=10,bold=row_index<=(2 if source_style else 1),color='FF4545' if str(value).startswith('-') else '20252C')
            if row_index==1 or row_index%2==0:cell.fill=PatternFill('solid',fgColor='F7FAFF')
    for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=20
    sheet.column_dimensions['B' if source_style else 'A'].width=34
    if source_style:
        sheet.column_dimensions['A'].width=8
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=sheet.max_column)
    content=BytesIO();book.save(content);book.close()
    return content.getvalue()


def export_enterprise_workbook(module, *, report='all', start='', end='', descending=True, hide_empty=True, unit='万元', decimals=2, scopes='', data_kinds='', window_years=0, trend_key=''):
    if window_years not in (0,3,5,10):raise ValueError('invalid_export_options')
    if unit not in UNIT_SCALES or not 0<=decimals<=6:raise ValueError('invalid_export_options')
    if any(value and not re.fullmatch(r'\d{4}',value) for value in (start,end)):raise ValueError('invalid_export_options')
    if start and end and start>end:raise ValueError('invalid_export_options')
    parsed=module.parsed_payload or {};metadata=parsed.get('metadata') or {}
    if metadata.get('unavailable') is True:raise ValueError('source_module_unavailable')
    periods=[_period(value) for value in parsed.get('periods') or []]
    rows=parsed.get('rows') or [];heads=parsed.get('head') or []
    if trend_key:
        statement_trend=getattr(module,'module_key','') in {'balance_sheet','income_statement','cash_flow_statement'}
        if (getattr(module,'module_key','') not in {'main_indicators','balance_sheet','income_statement','cash_flow_statement'}
                or not re.fullmatch(r'\d+' if statement_trend else r'\d+(?:_\d+)?',trend_key)
                or scopes not in {'合并期末','母公司期末'} or report not in {'annual','half','q1','q3'}
                or sum(isinstance(row,dict) and row.get('key')==trend_key for row in rows)!=1):
            raise ValueError('invalid_trend_options')
        hide_empty=False
        if statement_trend:
            return _workbook_bytes(module,_statement_trend_output(module,periods,rows,trend_key,report,start,end,window_years,scopes,unit,decimals),False,decimals)
    output=[];source_numbers=[];cell_formats={}
    impairment_record=bool(heads and isinstance(heads[0],list) and getattr(module,'module_key','') in {
        'receivables_impairment','other_receivables_impairment'})
    numbered_record=bool(heads and isinstance(heads[0],list) and getattr(module,'module_key','') in {
        'major_customers','major_suppliers','prepayments_top_five','receivables_impairment','other_receivables_impairment'})
    tag_columns=['companyTag'] if getattr(module,'module_key','')=='other_receivables_impairment' else ['companyTag','negativeTag']
    if heads and isinstance(heads[0],list):
        reports=metadata.get('report') or [];previous_header=None
        precise_record=metadata.get('precise_record') is True
        source_unit=(getattr(module,'request_params',{}) or {}).get('unit','')
        indices=list(range(len(heads)))
        if getattr(module,'module_key','') in {'receivables_top_five','other_receivables_top_five'} and reports:
            periods=[_period(value) for value in reports]
            latest=max(map(_order,periods),default=0)
            latest_quarter=(latest//10)*4+latest%10
            choices=set(report.split(','))
            indices=[index for index in indices if index<len(periods) and
                     (not start or periods[index][:4]>=start) and (not end or periods[index][:4]<=end) and
                     (not window_years or (_order(periods[index])//10)*4+_order(periods[index])%10>latest_quarter-window_years*4) and
                     ('all' in choices or _kind(periods[index]) in choices or 'latest' in choices and _order(periods[index])==latest)]
            indices.sort(key=lambda index:_order(periods[index]),reverse=descending)
        for index in indices:
            head=heads[index]
            columns=rows[index] if index<len(rows) else []
            header=[head[0],*[column[0] for column in columns]]
            if impairment_record:header.extend('企业类型标签' if key=='companyTag' else '负面信息标签' for key in tag_columns)
            if header!=previous_header:output.append(header);previous_header=header
            output.append([_period(reports[index]) if index<len(reports) else str(index+1)])
            tags=[]
            if impairment_record:
                for key in tag_columns:
                    groups=metadata.get(key) or []
                    group=groups[index] if index<len(groups) and isinstance(groups[index],list) and len(groups[index])==len(head) else []
                    tags.append(group)
            for i,name in enumerate(head[1:]):
                values=[]
                for column_index,column in enumerate(columns):
                    raw=column[i+1] if i+1<len(column) else None
                    if precise_record:
                        value=_amount(raw,'%' if re.search(r'[%％]',str(column[0])) else source_unit,unit,decimals)
                    elif impairment_record and _blank(raw):
                        value=None
                    elif impairment_record and (match:=re.fullmatch(r'(-?\d[\d,]*(?:\.\d+)?)(万|亿|元)',str(raw).strip())):
                        value=Decimal(match[1].replace(',',''))
                        cell_formats[(len(output)+2,column_index+3)]='###,###,##0.00"'+match[2]+'"'
                    else:value=raw
                    values.append(value)
                if impairment_record:
                    values.extend([' '.join(map(str,tag[i+1])).strip() if i+1<len(tag) and isinstance(tag[i+1],list) else None for tag in tags])
                output.append([name,*values])
    else:
        if heads and rows and all(isinstance(row,list) and re.fullmatch(r'\d{8}',str(row[0])) for row in rows):
            periods=[_period(row[0]) for row in rows]
            rows=[{'name':name,'values':[row[i+1] if i+1<len(row) else None for row in rows]} for i,name in enumerate(heads[1:])]
        if periods:
            reports=set(report.split(','));latest=max(map(_order,periods))
            column_types=next((row.get('values') or [] for row in rows if isinstance(row,dict) and row.get('key') in {'dataType','reportRange'}),[])
            indices=[i for i,p in enumerate(periods) if (not start or p[:4]>=start) and (not end or p[:4]<=end) and (
                'all' in reports or _kind(p) in reports or 'latest' in reports and _order(p)==latest or 'quarter' in reports and _kind(p) in {'q1','q3'})]
            indices.sort(key=lambda i:_order(periods[i]),reverse=descending)
            if window_years:
                quarter=lambda p:(_order(p)//10)*4+_order(p)%10
                cutoff=max(map(quarter,periods))-window_years*4
                indices=[i for i in indices if quarter(periods[i])>cutoff]
            indices=[i for i in indices if _column_selected(column_types[i] if i<len(column_types) else '',scopes,data_kinds)]
            output=[['报告期',*[periods[i] for i in indices]]]
            raw_data=(getattr(module,'raw_payload',{}) or {}).get('data') or {}
            export_headers=metadata.get('headExport') or []
            column_types=next((row.get('values') or [] for row in rows if isinstance(row,dict) and row.get('key')=='dataType'),[])
            for index,row in enumerate(_flatten(rows,raw_data.get('dataList') or [])):
                if trend_key and row.get('key') not in {trend_key,trend_key+'_2'}:continue
                source_unit=row.get('unit') or ''
                if not source_unit and getattr(module,'module_key','')=='main_business':
                    if re.search(r'[（(]%[）)]$',str(row.get('name',''))):source_unit='%'
                    elif re.match(r'^(120050|120062|900144|incomePrefix_|costPrefix_)',str(row.get('key',''))):
                        source_unit={'0':'元','3':'千元','4':'万元','6':'百万元','8':'亿元','9':'十亿元'}.get(str(module.request_params.get('unitCode')),'')
                if index+1<len(export_headers):
                    match=re.search(r'[（(]([^（）()]+)[）)]$',str(export_headers[index+1]))
                    if match:source_unit=match[1]
                name=str(row.get('name',row.get('label','')))
                if '每股' in name and source_unit=='元':source_unit='元/股'
                values=[row['values'][i] if i<len(row['values']) else None for i in indices]
                section=row.get('highlight') or row.get('level') in (0,'0')
                if hide_empty and not section and all(_blank(value) for value in values):continue
                label=name
                if source_unit:
                    display_unit=unit if source_unit in UNIT_SCALES else source_unit
                    suffix=re.search(r'[（(]([^（）()]+)[）)]$',label)
                    if suffix and suffix[1] in {*UNIT_SCALES,'%','倍','天','元/股'}:
                        label=label[:suffix.start()]+'（'+display_unit+'）'
                    else:label+='（'+display_unit+'）'
                converted=[_amount(value,'%' if source_unit in UNIT_SCALES and i<len(column_types) and (re.search(r'[%％]',str(column_types[i])) or column_types[i] in {'同比','占收入比'}) else source_unit,unit,decimals) for value,i in zip(values,indices)]
                if getattr(module,'module_key','')=='receivables_aging':
                    for position,value in enumerate(values):
                        match=re.fullmatch(r'(-?\d[\d,]*(?:\.\d+)?)(万|亿|元)',str(value).strip())
                        if match:
                            converted[position]=Decimal(match[1].replace(',',''))
                            cell_formats[(len(output)+2,position+3)]='###,###,##0.00"'+match[2]+'"'
                        elif _blank(value):converted[position]=None
                output.append([label,*converted])
                source_numbers.append(index+1)
        elif heads:output=[heads,*rows]
        else:output=[['企业预警通该栏目暂无数据' if metadata.get('empty') else '当前栏目数据尚待核对']]
    if trend_key:
        currency_values=next((row.get('values') or [] for row in rows if isinstance(row,dict) and row.get('key')=='displayCurrency'),[])
        currencies=[str(currency_values[i]) if i<len(currency_values) and not _blank(currency_values[i]) else '' for i in indices]
        unique_currencies=set(currencies)
        if len(unique_currencies)==1 and currencies[0]:
            for row in output[1:]:
                if row[0].endswith('（'+unit+'）'):
                    row[0]=row[0][:-1]+currencies[0]+'）'
        output=[['序号','报告期',*[row[0] for row in output[1:]]],
                *[[i,period,*[row[i] for row in output[1:]]] for i,period in enumerate(output[0][1:],1)]]
        if len(unique_currencies)>1:
            output[0].append('币种')
            for row,currency in zip(output[1:],currencies):row.append(currency)
    numbered_note=getattr(module,'module_key','') in {'cash_notes','inventory_notes','finance_costs','receivables_aging'} and bool(periods)
    source_style=not trend_key and bool(periods) and (module.category in {'indicators','statements'} or numbered_note)
    if numbered_record:
        source_style=True
        source_numbers=list(map(str,range(1,len(output))))
    return _workbook_bytes(module,output,source_style,decimals,
                           source_label=str(heads[0][0]) if numbered_record else str(heads[0]) if numbered_note and getattr(module,'module_key','')=='receivables_aging' else '项目名称' if numbered_note else '指标名称',
                           row_numbers=source_numbers if numbered_record else [str(number) for number in source_numbers] if numbered_note else None,
                           cell_formats=cell_formats)
